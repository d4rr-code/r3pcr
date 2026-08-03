"""Characterization tests for the supervisor analytics dashboard.

These lock the CURRENT context output of the ~540-line
``_analytics_context_response`` (rendered by the supervisor:dashboard view)
before it is refactored, so the section-by-section extraction can be proven
behavior-preserving. Assertions target the stable aggregate values (KPIs,
status/type/currency breakdowns, cost-by-type, feedback) computed from a small
deterministic dataset.

Run:  python manage.py test apps.supervisor --settings=config.settings_test
"""
from datetime import date, timedelta
from decimal import Decimal
import json

from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import User
from apps.shipments.models import HSCode, Shipment, ShipmentHSCode, StatusLog
from apps.computation.models import DutyComputation, ShipmentLineItem, ShippingAdvisory
from apps.computation.wmcda import calculate_ahp_weights
from apps.consignee.models import Feedback
from apps.supervisor.models import AuditLog, SystemConfig
from apps.supervisor.audit import log_audit


class SupervisorShipmentTrackingDisplayTests(TestCase):
    def setUp(self):
        self.supervisor = User.objects.create_user(
            username='sup_track', password='x', role='supervisor',
            email='sup_track@test.local', is_pending_approval=False,
        )
        self.declarant = User.objects.create_user(
            username='dec_track', password='x', role='declarant',
            email='dec_track@test.local', is_pending_approval=False,
        )
        self.consignee = User.objects.create_user(
            username='con_track', password='x', role='consignee',
            email='con_track@test.local', is_pending_approval=False,
        )
        self.shipment = Shipment.objects.create(
            hawb_number='R3PCR-TRACK-1',
            consignee=self.consignee,
            declarant=self.declarant,
            status='paid',
            shipment_type='lcl',
            job_order_reference='SRJJJ2511001234',
            container_number='TGHU1234567',
        )
        self.client.force_login(self.supervisor)

    def test_shipment_records_table_prioritizes_job_number(self):
        response = self.client.get(reverse('supervisor:shipment_records'))

        self.assertContains(response, '<th>Job Number</th>', html=False)
        self.assertNotContains(response, '<th>Import Type</th>', html=False)
        self.assertContains(response, 'SRJJJ2511001234')
        self.assertContains(response, 'TGHU1234567')

    def test_shipment_records_can_filter_by_mcda_recommendation(self):
        other = Shipment.objects.create(
            hawb_number='R3PCR-TRACK-2',
            consignee=self.consignee,
            declarant=self.declarant,
            status='paid',
            shipment_type='air',
        )
        ShippingAdvisory.objects.create(
            shipment=self.shipment, gross_weight=Decimal('1'),
            cargo_volume=Decimal('1'), declared_value=Decimal('1'),
            urgency_level='standard', distance_km=Decimal('2600'),
            lcl_score=Decimal('0.9'), fcl_score=Decimal('0.5'),
            air_score=Decimal('0.3'), recommended_type='lcl',
        )
        ShippingAdvisory.objects.create(
            shipment=other, gross_weight=Decimal('1'),
            cargo_volume=Decimal('1'), declared_value=Decimal('1'),
            urgency_level='standard', distance_km=Decimal('2600'),
            lcl_score=Decimal('0.3'), fcl_score=Decimal('0.5'),
            air_score=Decimal('0.9'), recommended_type='air',
        )

        response = self.client.get(reverse('supervisor:shipment_records'), {'mcda_rec': 'lcl'})

        self.assertContains(response, 'R3PCR-TRACK-1')
        self.assertNotContains(response, 'R3PCR-TRACK-2')
        self.assertContains(response, 'MCDA Recommendation')
        self.assertContains(response, 'mcda-lcl')

    def test_shipment_records_pagination_is_compact_for_many_pages(self):
        shipments = [
            Shipment(
                hawb_number=f'R3PCR-TRACK-PAGE-{idx:03d}',
                consignee=self.consignee,
                declarant=self.declarant,
                status='paid',
                shipment_type='lcl',
            )
            for idx in range(70)
        ]
        Shipment.objects.bulk_create(shipments)

        response = self.client.get(reverse('supervisor:shipment_records'), {'shipments_page': 8})
        links = response.context['shipments_page']['page_links']
        visible_numbers = [link['number'] for link in links if not link.get('ellipsis')]

        self.assertLess(len(links), response.context['shipments_page']['page_obj'].paginator.num_pages)
        self.assertIn({'ellipsis': True}, links)
        self.assertEqual(visible_numbers, [1, 6, 7, 8, 9, 10, 12])

    def test_supervisor_detail_shows_tracking_values_read_only(self):
        response = self.client.get(reverse('supervisor:shipment_detail', args=[self.shipment.id]))

        self.assertContains(response, 'Job Number')
        self.assertContains(response, 'SRJJJ2511001234')
        self.assertContains(response, 'TGHU1234567')


class SupervisorAuditTrailTests(TestCase):
    def setUp(self):
        self.supervisor = User.objects.create_user(
            username='sup_audit', password='x', role='supervisor',
            email='sup_audit@test.local', is_pending_approval=False,
        )
        self.client.force_login(self.supervisor)

    def test_report_download_is_logged_and_visible(self):
        response = self.client.get(reverse('supervisor:analytics_export'), {'format': 'xlsx'})

        self.assertEqual(response.status_code, 200)
        log = AuditLog.objects.get(action='report_download')
        self.assertEqual(log.user, self.supervisor)
        self.assertIn('analytics report', log.summary)

        page = self.client.get(reverse('supervisor:audit_trail'), {'action': 'report_download'})
        self.assertContains(page, 'Report Downloaded')
        self.assertContains(page, 'sup_audit')
        self.assertContains(page, 'analytics report')


class SupervisorIntelligenceTests(TestCase):
    def setUp(self):
        self.supervisor = User.objects.create_user(
            username='sup_intel', password='x', role='supervisor',
            email='sup_intel@test.local', is_pending_approval=False,
        )
        self.declarant = User.objects.create_user(
            username='dec_intel', password='x', role='declarant',
            email='dec_intel@test.local', is_pending_approval=False,
        )
        self.consignee = User.objects.create_user(
            username='con_intel', password='x', role='consignee',
            email='con_intel@test.local', is_pending_approval=False,
        )
        self.client.force_login(self.supervisor)

    def _shipment(self, hawb, status):
        return Shipment.objects.create(
            hawb_number=hawb,
            consignee=self.consignee,
            declarant=self.declarant,
            status=status,
            shipment_type='lcl',
        )

    def test_intelligence_page_reports_bottlenecks_risk_and_hs_review(self):
        base = timezone.now() - timedelta(days=6)
        shipment = self._shipment('R3PCR-INTEL-RISK', 'computed')
        Shipment.objects.filter(pk=shipment.pk).update(submitted_at=base)
        shipment.refresh_from_db()
        arrived = StatusLog.objects.create(
            shipment=shipment,
            changed_by=self.declarant,
            old_status='incoming',
            new_status='arrived',
        )
        computed = StatusLog.objects.create(
            shipment=shipment,
            changed_by=self.declarant,
            old_status='arrived',
            new_status='computed',
        )
        StatusLog.objects.filter(pk=arrived.pk).update(changed_at=base + timedelta(days=1))
        StatusLog.objects.filter(pk=computed.pk).update(changed_at=base + timedelta(days=2))

        completed = self._shipment('R3PCR-INTEL-DONE', 'billed')
        Shipment.objects.filter(pk=completed.pk).update(
            submitted_at=base,
            updated_at=base + timedelta(days=4),
        )

        hs = HSCode.objects.create(
            code='8471.60.90',
            description='Computer mouse and input units',
            duty_rate='1.00',
            chapter='84',
        )
        ShipmentHSCode.objects.create(shipment=shipment, hs_code=hs, is_confirmed=True)
        ShipmentLineItem.objects.create(
            shipment=shipment,
            description='wireless computer mouse',
            confidence='0.2000',
            source='ocr',
        )

        response = self.client.get(reverse('supervisor:intelligence'))

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'supervisor/intelligence.html')
        self.assertContains(response, 'Clearance Intelligence')
        self.assertTrue(response.context['stage_rows'])
        self.assertNotIn('computed', {row['status'] for row in response.context['stage_rows']})
        self.assertGreaterEqual(response.context['risk_distribution']['high'], 1)
        self.assertEqual(response.context['delay_model']['source'], 'Fallback rules')
        self.assertIn('projected_period_total', response.context['workload_forecast'])
        self.assertEqual(response.context['risk_filter'], 'high')
        self.assertEqual(response.context['hs_review']['historical_count'], 1)
        self.assertContains(response, '8471.60.90')
        self.assertContains(response, 'wireless computer mouse')
        self.assertContains(response, 'Projected Workload')
        self.assertContains(response, 'Delay Model Weights')
        year = timezone.localdate().year
        self.assertContains(response, f'?risk=all&forecast_unit=month&forecast_year={year}&forecast_model=sarima&forecast_months=3#delay-risk')
        self.assertContains(response, reverse('supervisor:intelligence_export') + f'?format=xlsx&risk=high&forecast_unit=month&forecast_year={year}&forecast_model=sarima')
        self.assertContains(response, reverse('supervisor:intelligence_export') + f'?format=pdf&risk=high&forecast_unit=month&forecast_year={year}&forecast_model=sarima')

    def test_intelligence_trains_delay_model_from_completed_shipments(self):
        today = timezone.localdate()
        for index in range(8):
            shipment = self._shipment(f'R3PCR-TRAIN-{index}', 'billed')
            submitted_at = timezone.now() - timedelta(days=8 + index)
            completed_at = submitted_at + timedelta(days=7 if index < 5 else 3)
            Shipment.objects.filter(pk=shipment.pk).update(
                submitted_at=submitted_at,
                updated_at=completed_at,
                has_deficiency=index < 5,
                estimated_arrival_date=today - timedelta(days=8 + index),
            )

        response = self.client.get(reverse('supervisor:intelligence'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['delay_model']['source'], 'Historical training')
        self.assertEqual(response.context['delay_model']['sample_count'], 8)
        self.assertGreater(response.context['delay_model']['weights']['deficiency'], 0)

    def test_intelligence_on_time_rate_uses_completed_kpi_result(self):
        submitted_at = timezone.now() - timedelta(days=20)
        on_time = self._shipment('R3PCR-KPI-ONTIME', 'billed')
        late = self._shipment('R3PCR-KPI-LATE', 'billed')
        Shipment.objects.filter(pk=on_time.pk).update(
            submitted_at=submitted_at,
            updated_at=submitted_at + timedelta(days=5),
        )
        Shipment.objects.filter(pk=late.pk).update(
            submitted_at=submitted_at,
            updated_at=submitted_at + timedelta(days=7),
        )

        response = self.client.get(reverse('supervisor:intelligence'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['completed_kpi_total'], 2)
        self.assertEqual(response.context['completed_kpi_on_time'], 1)
        self.assertEqual(response.context['completed_kpi_late'], 1)
        self.assertEqual(response.context['on_time_rate'], 50.0)
        self.assertContains(response, '1/2 completed on time')

    def test_intelligence_filters_delay_risk_rows(self):
        medium = self._shipment('R3PCR-RISK-MEDIUM', 'incoming')
        low = self._shipment('R3PCR-RISK-LOW', 'incoming')
        Shipment.objects.filter(pk=medium.pk).update(submitted_at=timezone.now() - timedelta(days=3))

        response = self.client.get(reverse('supervisor:intelligence'), {'risk': 'medium'})

        self.assertEqual(response.context['risk_filter'], 'medium')
        self.assertTrue(response.context['risk_rows'])
        self.assertTrue(all(row['label'] == 'Medium' for row in response.context['risk_rows']))
        self.assertContains(response, 'R3PCR-RISK-MEDIUM')
        self.assertNotContains(response, 'R3PCR-RISK-LOW')

    def test_intelligence_forecasts_workload_from_recent_incoming_trend(self):
        now = timezone.now()
        for index in range(4):
            shipment = self._shipment(f'R3PCR-FORECAST-RECENT-{index}', 'incoming')
            Shipment.objects.filter(pk=shipment.pk).update(submitted_at=now - timedelta(days=index))
        for index in range(2):
            shipment = self._shipment(f'R3PCR-FORECAST-PREV-{index}', 'incoming')
            Shipment.objects.filter(pk=shipment.pk).update(submitted_at=now - timedelta(days=35 + index))
        older = self._shipment('R3PCR-FORECAST-OLD', 'incoming')
        Shipment.objects.filter(pk=older.pk).update(submitted_at=now - timedelta(days=65))

        response = self.client.get(reverse('supervisor:intelligence'))

        forecast = response.context['workload_forecast']
        self.assertGreaterEqual(forecast['forecast_months'], 3)
        self.assertEqual(forecast['forecast_unit'], 'month')

        # Only 7 shipments exist, so SARIMA cannot fit. With the weighted moving
        # average removed there is no fallback: the page must declare the forecast
        # unavailable rather than report an empty projection as a real prediction.
        self.assertFalse(forecast['forecast_available'])
        self.assertIsNone(forecast['projected_period_total'])
        self.assertIsNone(forecast['projected_low'])
        self.assertIsNone(forecast['projected_high'])
        self.assertIsNone(forecast['trend_pct'])
        self.assertEqual(forecast['level'], '')
        self.assertNotEqual(forecast['level'], 'Light')
        self.assertEqual(forecast['confidence'], 'Unavailable')
        self.assertEqual(forecast['model_source'], 'Forecast unavailable')
        self.assertTrue(forecast['unavailable_reason'])
        self.assertContains(response, 'Insufficient data for forecast')

        # SARIMA is the only registered model.
        self.assertEqual(
            {row['key'] for row in forecast['model_comparison']},
            {'sarima'},
        )
        self.assertEqual(forecast['forecast_model'], 'sarima')
        self.assertEqual(forecast['recommended_model']['key'], 'sarima')
        self.assertEqual(forecast['displayed_model']['key'], 'sarima')
        expected_months = timezone.localdate().month + 3
        self.assertEqual(forecast['forecast_months'], 3)
        self.assertEqual(len(forecast['chart']['labels']), expected_months)
        self.assertEqual(len(forecast['chart']['historical_values']), expected_months)
        self.assertEqual(len(forecast['chart']['forecast_values']), expected_months)
        self.assertEqual(forecast['chart']['historical_label'], 'Historical monthly volume')
        self.assertEqual(forecast['chart']['historical_values'][-3:], [None, None, None])
        # With no forecast produced, the period table lists only the actual months —
        # no projected rows are appended.
        self.assertEqual(len(forecast['period_rows']), timezone.localdate().month)
        self.assertTrue(all(row['date_range'] == 'Actual' for row in forecast['period_rows']))
        self.assertContains(response, 'SARIMA')
        self.assertContains(response, 'fallbackLineChart')
        self.assertNotContains(response, 'Weighted Moving Average')

        for row in forecast['model_comparison']:
            self.assertIn('mae', row)
            self.assertIn('rmse', row)
            self.assertIn('mape', row)
            self.assertIn('rmse_label', row)
            self.assertIn('mape_label', row)
            self.assertEqual(row['key'], 'sarima')
            self.assertIn(row['status'], ('Active', 'Unavailable'))

        year_override = self.client.get(reverse('supervisor:intelligence'), {
            'forecast_unit': 'year',
        })
        self.assertEqual(year_override.context['workload_forecast']['forecast_unit'], 'month')
        self.assertEqual(
            year_override.context['workload_forecast']['chart']['historical_label'],
            'Historical monthly volume',
        )
        self.assertIn('Next 3 months', year_override.context['workload_forecast']['forecast_label'])

    def test_intelligence_exports_xlsx_and_pdf(self):
        self._shipment('R3PCR-INTEL-EXPORT', 'billed')

        xlsx = self.client.get(reverse('supervisor:intelligence_export'), {'format': 'xlsx'})
        self.assertEqual(xlsx.status_code, 200)
        self.assertIn('spreadsheet', xlsx['Content-Type'])
        self.assertTrue(xlsx.content.startswith(b'PK'))

        pdf = self.client.get(reverse('supervisor:intelligence_export'), {'format': 'pdf'})
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf['Content-Type'], 'application/pdf')
        self.assertTrue(pdf.content.startswith(b'%PDF'))


class AuditTrailExportTests(TestCase):
    """CSV download on the supervisor audit trail page."""

    def setUp(self):
        self.supervisor = User.objects.create_user(
            username='sup_audit_csv', password='pw', role='supervisor',
            email='sup.audit@example.com', first_name='Ana', last_name='Reyes',
        )
        self.client.force_login(self.supervisor)
        log_audit('login', 'Supervisor signed in.', user=self.supervisor)
        log_audit('config_update', 'Updated MCDA weights.', user=self.supervisor)

    def _download(self, **params):
        return self.client.get(reverse('supervisor:audit_trail_export'), params)

    def test_csv_download_returns_rows_with_headers(self):
        response = self._download()

        self.assertEqual(response.status_code, 200)
        self.assertIn('text/csv', response['Content-Type'])
        self.assertIn('attachment;', response['Content-Disposition'])
        self.assertIn('R3PCR_Audit_Trail_', response['Content-Disposition'])
        self.assertTrue(response['Content-Disposition'].rstrip('"').endswith('.csv'))

        body = response.content.decode('utf-8-sig')
        lines = [line for line in body.splitlines() if line.strip()]

        # Metadata block sits above the table.
        self.assertEqual(lines[0], 'R3-PCR Audit Trail Report')
        self.assertIn(
            'Date/Time,Action,User,Email,Role,Shipment,Summary,IP Address',
            body,
        )
        self.assertIn('Updated MCDA weights.', body)
        self.assertIn('Ana Reyes', body)
        self.assertIn('sup.audit@example.com', body)

    def test_csv_includes_report_metadata_block(self):
        response = self._download(period='this_week')
        body = response.content.decode('utf-8-sig')

        self.assertIn('R3-PCR Audit Trail Report', body)
        self.assertIn('Generated by:,Ana Reyes', body)
        self.assertIn('Generated on:', body)
        self.assertIn('Reporting Period:,This week', body)
        self.assertIn('Applied Filters:', body)
        self.assertIn('Total Records:', body)

    def test_metadata_reflects_applied_filters(self):
        body = self._download(action='config_update').content.decode('utf-8-sig')

        self.assertIn('Action: System Configuration Updated', body)
        self.assertIn('User: All', body)
        self.assertIn('Shipment: All', body)

    def test_csv_starts_with_bom_for_excel(self):
        self.assertTrue(self._download().content.startswith(b'\xef\xbb\xbf'))

    def test_export_respects_page_filters(self):
        response = self._download(action='config_update')

        body = response.content.decode('utf-8-sig')
        self.assertIn('Updated MCDA weights.', body)
        self.assertNotIn('Supervisor signed in.', body)

    def test_formula_injection_is_neutralised(self):
        log_audit('config_update', '=cmd|calc!A1', user=self.supervisor)

        body = self._download(action='config_update').content.decode('utf-8-sig')

        self.assertIn("'=cmd|calc!A1", body)

    def test_download_is_itself_audited(self):
        self._download()

        entry = AuditLog.objects.filter(action='report_download').first()
        self.assertIsNotNone(entry)
        self.assertEqual(entry.details['format'], 'csv')

    def test_non_supervisor_cannot_download(self):
        self.client.logout()
        consignee = User.objects.create_user(
            username='con_audit_csv', password='pw', role='consignee',
        )
        self.client.force_login(consignee)

        response = self._download()

        self.assertNotEqual(response.status_code, 200)


class DashboardReportCsvTests(TestCase):
    """CSV format for the analytics and clearance intelligence reports."""

    def setUp(self):
        self.supervisor = User.objects.create_user(
            username='sup_csv_reports', password='pw', role='supervisor',
        )
        self.client.force_login(self.supervisor)
        SystemConfig.set('exchange_rates_last_success', timezone.localdate().isoformat())
        SystemConfig.set('exchange_rates_last_attempt', timezone.localdate().isoformat())

    def test_analytics_csv_export(self):
        response = self.client.get(reverse('supervisor:analytics_export'), {'format': 'csv'})

        self.assertEqual(response.status_code, 200)
        self.assertIn('text/csv', response['Content-Type'])
        self.assertIn('R3PCR_Analytics_Report_', response['Content-Disposition'])
        body = response.content.decode('utf-8-sig')
        self.assertIn('R3-PCR Analytics Report', body)
        self.assertIn('Executive Summary', body)
        self.assertIn('Status Pipeline', body)

    def test_intelligence_csv_export(self):
        response = self.client.get(reverse('supervisor:intelligence_export'), {'format': 'csv'})

        self.assertEqual(response.status_code, 200)
        self.assertIn('text/csv', response['Content-Type'])
        body = response.content.decode('utf-8-sig')
        self.assertIn('R3-PCR Clearance Intelligence Report', body)
        self.assertIn('Executive Summary', body)
        self.assertIn('Delay Risk', body)

    def test_existing_formats_are_unaffected(self):
        xlsx = self.client.get(reverse('supervisor:analytics_export'), {'format': 'xlsx'})
        self.assertTrue(xlsx.content.startswith(b'PK'))

        pdf = self.client.get(reverse('supervisor:analytics_export'), {'format': 'pdf'})
        self.assertTrue(pdf.content.startswith(b'%PDF'))

    def test_generate_reports_card_renders(self):
        response = self.client.get(reverse('supervisor:dashboard'))

        self.assertContains(response, 'Generate Reports')
        self.assertContains(response, 'Reporting Period')
        self.assertContains(response, 'Apply Filters')
        self.assertContains(response, 'Export CSV')
        self.assertContains(response, 'Export PDF')
        # Excel was dropped from the card as redundant with CSV; the backend
        # still serves format=xlsx for anyone hitting the URL directly.
        self.assertNotContains(response, 'Export Excel')
        self.assertContains(response, 'Last Month')
        self.assertContains(response, 'This Quarter')


class ReportingPeriodTests(SimpleTestCase):
    """Reporting-period preset resolution shared by the operational exports."""

    def _resolve(self, **params):
        from django.test import RequestFactory
        from apps.supervisor.views.reports import resolve_report_period

        return resolve_report_period(RequestFactory().get('/x/', params))

    def test_today_preset_spans_a_single_day(self):
        today = timezone.localdate().isoformat()
        period = self._resolve(period='today')

        self.assertEqual(period['date_from'], today)
        self.assertEqual(period['date_to'], today)
        self.assertEqual(period['label'], 'Today')

    def test_this_month_starts_on_the_first(self):
        period = self._resolve(period='this_month')

        self.assertEqual(period['date_from'], timezone.localdate().replace(day=1).isoformat())
        self.assertEqual(period['date_to'], timezone.localdate().isoformat())

    def test_last_month_ends_before_this_month_starts(self):
        period = self._resolve(period='last_month')
        first_this_month = timezone.localdate().replace(day=1)

        self.assertLess(date.fromisoformat(period['date_to']), first_this_month)
        self.assertEqual(date.fromisoformat(period['date_from']).day, 1)

    def test_preset_overrides_raw_dates(self):
        period = self._resolve(period='today', date_from='2020-01-01', date_to='2020-12-31')

        self.assertEqual(period['date_from'], timezone.localdate().isoformat())

    def test_raw_dates_are_preserved_without_a_preset(self):
        period = self._resolve(date_from='2024-01-01', date_to='2024-06-30')

        self.assertEqual(period['date_from'], '2024-01-01')
        self.assertEqual(period['date_to'], '2024-06-30')
        self.assertEqual(period['period'], 'custom')

    def test_no_filters_means_all_time(self):
        period = self._resolve()

        self.assertEqual(period['date_from'], '')
        self.assertEqual(period['date_to'], '')
        self.assertEqual(period['label'], 'All time')

    def test_malformed_inputs_are_ignored(self):
        period = self._resolve(period='not_a_period', date_from='garbage')

        self.assertEqual(period['period'], '')
        self.assertEqual(period['date_from'], '')


class ShipmentRegisterExportTests(TestCase):
    """Supervisor shipment register export."""

    def setUp(self):
        self.supervisor = User.objects.create_user(
            username='sup_ship_rep', password='pw', role='supervisor',
        )
        self.consignee = User.objects.create_user(
            username='con_ship_rep', password='pw', role='consignee',
            first_name='Rosa', last_name='Lim', company_name='Lim Trading',
        )
        self.client.force_login(self.supervisor)
        self.recent = Shipment.objects.create(
            hawb_number='R3PCR-REP-RECENT', consignee=self.consignee,
            status='computed', shipment_type='air', urgency='urgent',
            import_type='commercial', declared_value=Decimal('1500.00'),
            invoice_currency='USD',
        )
        self.old = Shipment.objects.create(
            hawb_number='R3PCR-REP-OLD', consignee=self.consignee,
            status='billed', shipment_type='fcl',
        )
        Shipment.objects.filter(pk=self.old.pk).update(
            submitted_at=timezone.now() - timedelta(days=400)
        )

    def _export(self, **params):
        return self.client.get(reverse('supervisor:shipment_records_export'), params)

    def test_csv_export_contains_expected_headers_and_rows(self):
        response = self._export()

        self.assertEqual(response.status_code, 200)
        self.assertIn('text/csv', response['Content-Type'])
        body = response.content.decode('utf-8-sig')
        self.assertIn(
            'HAWB,Consignee,Declarant,Type,Status,Urgency,Import Type,'
            'Declared Value,Currency,Submitted,Last Updated,KPI Status',
            body,
        )
        self.assertIn('R3PCR-REP-RECENT', body)
        self.assertIn('Lim Trading', body)
        self.assertIn('Unassigned', body)

    def test_pdf_export_returns_a_pdf(self):
        response = self._export(format='pdf')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))

    def test_status_filter_is_honoured(self):
        body = self._export(status_f='computed').content.decode('utf-8-sig')

        self.assertIn('R3PCR-REP-RECENT', body)
        self.assertNotIn('R3PCR-REP-OLD', body)

    def test_reporting_period_excludes_older_shipments(self):
        body = self._export(period='this_month').content.decode('utf-8-sig')

        self.assertIn('R3PCR-REP-RECENT', body)
        self.assertNotIn('R3PCR-REP-OLD', body)

    def test_export_is_audited(self):
        self._export()

        entry = AuditLog.objects.filter(action='report_download').first()
        self.assertIsNotNone(entry)
        self.assertEqual(entry.details['format'], 'csv')

    def test_non_supervisor_cannot_export(self):
        self.client.logout()
        self.client.force_login(self.consignee)

        self.assertNotEqual(self._export().status_code, 200)


class UserAccountsExportTests(TestCase):
    """Supervisor user accounts export."""

    def setUp(self):
        self.supervisor = User.objects.create_user(
            username='sup_user_rep', password='pw', role='supervisor',
        )
        self.client.force_login(self.supervisor)
        User.objects.create_user(
            username='con_active_rep', password='pw', role='consignee',
            first_name='Ana', last_name='Cruz', email='ana@example.com',
            company_name='Cruz Imports', is_active=True,
        )
        inactive = User.objects.create_user(
            username='dec_inactive_rep', password='pw', role='declarant',
            first_name='Ben', last_name='Tan',
        )
        User.objects.filter(pk=inactive.pk).update(is_active=False)

    def _export(self, **params):
        return self.client.get(reverse('supervisor:users_export'), params)

    def test_csv_export_contains_headers_and_accounts(self):
        response = self._export()

        self.assertEqual(response.status_code, 200)
        body = response.content.decode('utf-8-sig')
        self.assertIn(
            'Username,Full Name,Email,Role,Company,Phone,Active,'
            'Email Verified,Date Joined',
            body,
        )
        self.assertIn('con_active_rep', body)
        self.assertIn('Cruz Imports', body)

    def test_pdf_export_returns_a_pdf(self):
        response = self._export(format='pdf')

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))

    def test_role_filter_is_honoured(self):
        body = self._export(role='declarant').content.decode('utf-8-sig')

        self.assertIn('dec_inactive_rep', body)
        self.assertNotIn('con_active_rep', body)

    def test_status_filter_is_honoured(self):
        body = self._export(status='inactive').content.decode('utf-8-sig')

        self.assertIn('dec_inactive_rep', body)
        self.assertNotIn('con_active_rep', body)

    def test_export_is_audited(self):
        self._export()

        self.assertTrue(AuditLog.objects.filter(action='report_download').exists())

    def test_non_supervisor_cannot_export(self):
        self.client.logout()
        consignee = User.objects.create_user(
            username='con_blocked_rep', password='pw', role='consignee',
        )
        self.client.force_login(consignee)

        self.assertNotEqual(self._export().status_code, 200)


class ForecastModelTests(SimpleTestCase):
    """SARIMA-only forecasting: model registry, metrics, and unavailable state."""

    def test_sarima_is_the_only_registered_model(self):
        from apps.supervisor.views.intelligence import FORECAST_MODELS

        self.assertEqual([model['key'] for model in FORECAST_MODELS], ['sarima'])

    def test_backtest_split_is_28_8_over_the_36_month_window(self):
        from apps.supervisor.views.intelligence import _backtest_split

        # Integer truncation makes the realised ratio 77.8%/22.2%, not 80/20.
        self.assertEqual(_backtest_split(36), (28, 8))
        self.assertEqual(_backtest_split(5), (0, 0))  # below the 6-observation floor

    def test_sufficient_history_produces_sarima_forecast_with_all_metrics(self):
        from apps.supervisor.views.intelligence import _forecast_model_comparison

        # 36 months of seasonal-ish volume, comfortably past SARIMA's 24-point floor.
        counts = [10, 11, 14, 13, 14, 10, 8, 14, 12, 13, 15, 13,
                  10, 13, 12, 18, 11, 9, 8, 18, 9, 18, 10, 14,
                  11, 12, 16, 19, 13, 12, 6, 16, 14, 16, 16, 15]

        forecasts, source, recommended, rows = _forecast_model_comparison(counts, 3, 'month')

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['key'], 'sarima')
        self.assertEqual(rows[0]['status'], 'Active')
        self.assertEqual(len(forecasts), 3)
        self.assertEqual(recommended['key'], 'sarima')
        self.assertIn('SARIMA', source)
        for metric in ('mae', 'rmse', 'mape'):
            self.assertIsNotNone(recommended[metric], f'{metric} should be reported')

    def test_insufficient_history_yields_no_forecast_and_no_fallback(self):
        from apps.supervisor.views.intelligence import _forecast_model_comparison

        counts = [1, 0, 2, 1, 0, 1, 1, 0]  # sum well below SARIMA's minimum

        forecasts, source, recommended, rows = _forecast_model_comparison(counts, 3, 'month')

        self.assertEqual(forecasts, [])
        self.assertEqual(source, 'Forecast unavailable')
        self.assertEqual(recommended['key'], 'sarima')
        self.assertIsNone(recommended['mae'])
        self.assertTrue(recommended['unavailable_reason'])
        self.assertEqual([row['key'] for row in rows], ['sarima'])
        self.assertEqual(rows[0]['status'], 'Unavailable')


class WmcdaAhpTests(TestCase):
    def test_ahp_weights_sum_to_100_and_return_consistency_ratio(self):
        result = calculate_ahp_weights({
            'ahp_cost_time': '3',
            'ahp_cost_weight': '5',
            'ahp_cost_distance': '7',
            'ahp_time_weight': '3',
            'ahp_time_distance': '5',
            'ahp_weight_distance': '3',
        })

        self.assertEqual(sum(result['weights_pct'].values()), 100)
        self.assertIn('cost', result['weights_pct'])
        self.assertGreaterEqual(result['consistency_ratio'], 0)

    def test_config_wmcda_can_apply_ahp_weights(self):
        supervisor = User.objects.create_user(
            username='sup_ahp', password='x', role='supervisor',
            email='supahp@test.local', is_pending_approval=False,
        )
        self.client.force_login(supervisor)
        resp = self.client.post(reverse('supervisor:config_wmcda'), {
            'action': 'apply_ahp',
            'ahp_cost_time': '3',
            'ahp_cost_weight': '5',
            'ahp_cost_distance': '7',
            'ahp_time_weight': '3',
            'ahp_time_distance': '5',
            'ahp_weight_distance': '3',
        })

        self.assertRedirects(resp, reverse('supervisor:config_wmcda'), fetch_redirect_response=False)
        weights = [
            int(SystemConfig.get('wmcda_w_cost', '0')),
            int(SystemConfig.get('wmcda_w_time', '0')),
            int(SystemConfig.get('wmcda_w_weight', '0')),
            int(SystemConfig.get('wmcda_w_distance', '0')),
        ]
        self.assertEqual(sum(weights), 100)
        self.assertEqual(SystemConfig.get('wmcda_weight_method', ''), 'saaty_ahp')
        self.assertTrue(SystemConfig.get('wmcda_ahp_matrix', ''))

    def test_config_wmcda_page_renders_ahp_controls(self):
        supervisor = User.objects.create_user(
            username='sup_ahp_get', password='x', role='supervisor',
            email='supahpget@test.local', is_pending_approval=False,
        )
        self.client.force_login(supervisor)

        resp = self.client.get(reverse('supervisor:config_wmcda'))

        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Saaty AHP Weight Derivation')
        self.assertContains(resp, 'ahp_cost_time')


class AnalyticsDashboardContextTests(TestCase):
    def setUp(self):
        self.supervisor = User.objects.create_user(
            username='sup', password='x', role='supervisor',
            email='sup@test.local', is_pending_approval=False,
        )
        self.dec1 = User.objects.create_user(
            username='dec1', password='x', role='declarant',
            email='dec1@test.local', is_pending_approval=False,
        )
        self.dec2 = User.objects.create_user(
            username='dec2', password='x', role='declarant',
            email='dec2@test.local', is_pending_approval=False,
        )
        self.consignee = User.objects.create_user(
            username='con1', password='x', role='consignee',
            email='con1@test.local', is_pending_approval=False,
        )

        def _ship(hawb, status, stype, currency, declarant):
            return Shipment.objects.create(
                hawb_number=hawb, consignee=self.consignee, declarant=declarant,
                status=status, shipment_type=stype, invoice_currency=currency,
            )

        # total_all=5 ; status: incoming/arrived/computed/approved/billed = 1 each
        # type: air=2, lcl=2, fcl=1 ; currency: USD=3, EUR=1, JPY=1
        self.s1 = _ship('A-1', 'incoming', 'air', 'USD', self.dec1)
        self.s2 = _ship('A-2', 'arrived',  'lcl', 'USD', self.dec1)
        self.s3 = _ship('A-3', 'computed', 'fcl', 'EUR', self.dec2)
        self.s4 = _ship('A-4', 'approved', 'air', 'USD', self.dec2)
        self.s5 = _ship('A-5', 'billed',   'lcl', 'JPY', self.dec1)

        # Cost-by-type: lcl avg 1500 (count 2), air avg 500 (count 1), fcl none
        DutyComputation.objects.create(shipment=self.s1, total_landed_cost=Decimal('500'))
        DutyComputation.objects.create(shipment=self.s2, total_landed_cost=Decimal('1000'))
        DutyComputation.objects.create(shipment=self.s5, total_landed_cost=Decimal('2000'))

        # Feedback: ratings 5,4,2 -> total 3, avg 3.7, positive (>=4) = 2
        Feedback.objects.create(consignee=self.consignee, shipment=self.s1, rating=5, comment='a')
        Feedback.objects.create(consignee=self.consignee, shipment=self.s2, rating=4, comment='b')
        Feedback.objects.create(consignee=self.consignee, shipment=self.s3, rating=2, comment='c')

        self.client.force_login(self.supervisor)
        self.url = reverse('supervisor:dashboard')

    def _ctx(self):
        resp = self.client.get(self.url)
        self.assertEqual(resp.status_code, 200)
        self.assertTemplateUsed(resp, 'supervisor/analytics.html')
        return resp.context

    def test_kpi_strip_counts(self):
        ctx = self._ctx()
        self.assertEqual(ctx['total_all'], 5)
        self.assertEqual(ctx['total_incoming'], 1)
        self.assertEqual(ctx['total_arrived'], 1)
        self.assertEqual(ctx['total_computed'], 1)
        self.assertEqual(ctx['total_approved'], 1)
        self.assertEqual(ctx['total_rejected'], 0)
        self.assertEqual(ctx['total_declarants'], 2)
        self.assertEqual(ctx['total_consignees'], 1)

    def test_shipment_type_counts(self):
        ctx = self._ctx()
        self.assertEqual(ctx['shipment_type_counts'], {'air': 2, 'lcl': 2, 'fcl': 1})

    def test_currency_breakdown(self):
        ctx = self._ctx()
        self.assertEqual(ctx['currency_total'], 5)
        by_code = {r['code']: r['count'] for r in ctx['currency_breakdown']}
        self.assertEqual(by_code['USD'], 3)
        self.assertEqual(by_code['EUR'], 1)
        self.assertEqual(by_code['JPY'], 1)

    def test_cost_by_type(self):
        ctx = self._ctx()
        by_code = {r['code']: r for r in ctx['cost_by_type']}
        self.assertEqual(by_code['lcl']['count'], 2)
        self.assertEqual(by_code['lcl']['avg'], 1500.0)
        self.assertEqual(by_code['lcl']['total'], 3000.0)
        self.assertEqual(by_code['air']['count'], 1)
        self.assertEqual(by_code['air']['avg'], 500.0)
        self.assertEqual(by_code['fcl']['count'], 0)

    def test_feedback_summary(self):
        ctx = self._ctx()
        fb = ctx['feedback_summary']
        self.assertEqual(fb['total'], 3)
        self.assertEqual(fb['avg_rating'], 3.7)
        self.assertEqual(fb['positive'], 2)

    def test_status_rows_reflect_seeded_statuses(self):
        ctx = self._ctx()
        counts = {r['key']: r['count'] for r in ctx['status_rows']}
        self.assertEqual(counts.get('incoming'), 1)
        self.assertEqual(counts.get('arrived'), 1)
        self.assertEqual(counts.get('computed'), 1)
        self.assertEqual(counts.get('approved'), 1)
        self.assertEqual(counts.get('billed'), 1)
        self.assertEqual(ctx['chart_total'], 5)

    def test_declarant_filter_narrows_chart_total(self):
        resp = self.client.get(self.url, {'declarant': 'dec1'})
        self.assertEqual(resp.status_code, 200)
        # dec1 owns s1, s2, s5 -> 3
        self.assertEqual(resp.context['chart_total'], 3)

    def test_date_filter_updates_dashboard_kpis_and_type_counts(self):
        today = timezone.localdate()
        old = timezone.now() - timedelta(days=7)
        Shipment.objects.filter(pk__in=[self.s2.pk, self.s3.pk, self.s4.pk, self.s5.pk]).update(submitted_at=old)
        Shipment.objects.filter(pk=self.s1.pk).update(submitted_at=timezone.now())

        resp = self.client.get(self.url, {
            'date_from': today.isoformat(),
            'date_to': today.isoformat(),
        })

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context['total_all'], 1)
        self.assertEqual(resp.context['chart_total'], 1)
        self.assertEqual(resp.context['total_incoming'], 1)
        self.assertEqual(resp.context['total_arrived'], 0)
        self.assertEqual(resp.context['shipment_type_counts'], {'air': 1, 'lcl': 0, 'fcl': 0})
        self.assertEqual(resp.context['feedback_summary']['total'], 1)
        self.assertEqual(resp.context['feedback_summary']['avg_rating'], 5.0)
        monthly_data = json.loads(resp.context['monthly_chart_data'])
        self.assertEqual(sum(value or 0 for value in monthly_data), 1)
        self.assertEqual(len(json.loads(resp.context['monthly_chart_labels'])), 12)

    def test_all_years_overview_groups_chart_by_year(self):
        old_year = timezone.now().replace(year=timezone.localdate().year - 1)
        Shipment.objects.filter(pk__in=[self.s1.pk, self.s2.pk]).update(submitted_at=old_year)

        resp = self.client.get(self.url, {'overview_range': 'all'})

        self.assertEqual(resp.status_code, 200)
        labels = json.loads(resp.context['monthly_chart_labels'])
        data = json.loads(resp.context['monthly_chart_data'])
        self.assertIn(str(timezone.localdate().year - 1), labels)
        self.assertIn(str(timezone.localdate().year), labels)
        self.assertEqual(sum(data), 5)
        self.assertTrue(resp.context['monthly_chart_has_data'])
        self.assertContains(resp, 'All Years')

    def test_date_filter_controls_render_range_presets(self):
        """The Generate Reports card exposes the reporting-period presets.

        Replaces the earlier 'This Month / Last 30 Days / Custom Range' button
        row; 'Last 30 Days' is intentionally no longer offered.
        """
        resp = self.client.get(self.url)

        self.assertContains(resp, 'Reporting Period')
        for option in ('All Time', 'Today', 'This Week', 'This Month',
                       'Last Month', 'This Quarter', 'This Year', 'Custom Range'):
            self.assertContains(resp, option)
        self.assertNotContains(resp, 'Last 30 Days')

    def test_live_status_counts_respect_date_and_declarant_filters(self):
        today = timezone.localdate()
        old = timezone.now() - timedelta(days=7)
        Shipment.objects.filter(pk__in=[self.s2.pk, self.s3.pk, self.s4.pk, self.s5.pk]).update(submitted_at=old)
        Shipment.objects.filter(pk=self.s1.pk).update(submitted_at=timezone.now())

        resp = self.client.get(reverse('supervisor:analytics_status_counts'), {
            'date_from': today.isoformat(),
            'date_to': today.isoformat(),
            'declarant': 'dec1',
        })

        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data['total'], 1)
        self.assertEqual(data['counts']['incoming']['count'], 1)
        self.assertEqual(data['counts']['arrived']['count'], 0)

    def test_wmcda_scoreboard_and_agreement(self):
        def _adv(shipment, recommended):
            return ShippingAdvisory.objects.create(
                shipment=shipment, gross_weight=Decimal('1'),
                cargo_volume=Decimal('1'), declared_value=Decimal('1'),
                urgency_level='standard', distance_km=Decimal('2600'),
                lcl_score=Decimal('0.9'), fcl_score=Decimal('0.5'),
                air_score=Decimal('0.3'), recommended_type=recommended,
            )
        _adv(self.s2, 'lcl')   # declared lcl -> rec lcl (match)
        _adv(self.s5, 'lcl')   # declared lcl -> rec lcl (match)
        _adv(self.s1, 'air')   # declared air -> rec air (match)

        ctx = self._ctx()
        self.assertEqual(ctx['wmcda_total'], 3)
        board = {r['key']: r['count'] for r in ctx['wmcda_scoreboard']}
        self.assertEqual(board['lcl'], 2)
        self.assertEqual(board['air'], 1)
        self.assertEqual(board['fcl'], 0)
        # all three declared==recommended -> 100% agreement
        self.assertEqual(ctx['wmcda_comparison_agreement'], 100)

    def test_wmcda_drilldown_links_filter_by_recommendation(self):
        ShippingAdvisory.objects.create(
            shipment=self.s2, gross_weight=Decimal('1'),
            cargo_volume=Decimal('1'), declared_value=Decimal('1'),
            urgency_level='standard', distance_km=Decimal('2600'),
            lcl_score=Decimal('0.9'), fcl_score=Decimal('0.5'),
            air_score=Decimal('0.3'), recommended_type='lcl',
        )

        response = self.client.get(self.url)

        self.assertContains(response, '?mcda_rec=lcl')

    def test_declarant_performance_speed_and_volume(self):
        base = timezone.now().replace(microsecond=0)
        arr = StatusLog.objects.create(
            shipment=self.s2, new_status='arrived', old_status='incoming',
            changed_by=self.dec1,
        )
        StatusLog.objects.filter(pk=arr.pk).update(changed_at=base)
        comp = StatusLog.objects.create(
            shipment=self.s2, new_status='computed', old_status='arrived',
            changed_by=self.dec1,
        )
        StatusLog.objects.filter(pk=comp.pk).update(changed_at=base + timedelta(hours=2))

        ctx = self._ctx()
        by_user = {d['username']: d for d in ctx['declarant_data']}
        self.assertEqual(by_user['dec1']['total_processed'], 1)
        self.assertEqual(by_user['dec1']['avg_hours'], 2.0)
        self.assertEqual(by_user['dec2']['total_processed'], 0)

    def test_non_supervisor_is_redirected(self):
        self.client.force_login(self.dec1)
        resp = self.client.get(self.url)
        self.assertEqual(resp.status_code, 302)


class AnalyticsExportTests(TestCase):
    """Smoke-tests for the analytics report download (PDF + XLSX) so the export
    can't silently break while we adjust the PDF layout."""
    def setUp(self):
        self.supervisor = User.objects.create_user(
            username='sup_ex', password='x', role='supervisor',
            email='supex@test.local', is_pending_approval=False)
        self.consignee = User.objects.create_user(
            username='con_ex', password='x', role='consignee',
            email='conex@test.local', is_pending_approval=False)
        Shipment.objects.create(
            hawb_number='EX-1', consignee=self.consignee, status='billed',
            shipment_type='lcl', invoice_currency='USD')
        self.client.force_login(self.supervisor)
        self.url = reverse('supervisor:analytics_export')

    def test_pdf_export(self):
        resp = self.client.get(self.url, {'format': 'pdf'})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertTrue(resp.content.startswith(b'%PDF'))
        self.assertIn('.pdf', resp['Content-Disposition'])

    def test_xlsx_export(self):
        resp = self.client.get(self.url, {'format': 'xlsx'})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheet', resp['Content-Type'])
        self.assertGreater(len(resp.content), 0)
        from io import BytesIO
        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(resp.content), read_only=True)
        self.assertIn('Incoming Workload Forecast', wb.sheetnames)
        self.assertIn('Forecast Periods', wb.sheetnames)
        self.assertIn('Forecast Model Comparison', wb.sheetnames)
        summary_values = [row[0] for row in wb['Summary'].iter_rows(values_only=True) if row and row[0]]
        self.assertIn('Projected Incoming Workload', summary_values)

    def test_export_requires_supervisor(self):
        self.client.logout()
        other = User.objects.create_user(
            username='con_ey', password='x', role='consignee',
            email='coney@test.local', is_pending_approval=False)
        self.client.force_login(other)
        resp = self.client.get(self.url, {'format': 'pdf'})
        self.assertEqual(resp.status_code, 302)


class UserManagementTests(TestCase):
    def setUp(self):
        self.supervisor = User.objects.create_user(
            username='sup_users', password='x', role='supervisor',
            email='supusers@test.local', is_pending_approval=False)
        self.pending = User.objects.create_user(
            username='pending_con', password='x', role='consignee',
            email='pending@test.local', is_active=False,
            is_pending_approval=True, email_verified=False)
        self.client.force_login(self.supervisor)

    def test_cannot_approve_unverified_email_registration(self):
        resp = self.client.post(reverse('supervisor:approve_registration', args=[self.pending.id]))

        self.assertRedirects(resp, reverse('supervisor:users'), fetch_redirect_response=False)
        self.pending.refresh_from_db()
        self.assertFalse(self.pending.is_active)
        self.assertTrue(self.pending.is_pending_approval)

    def test_unverified_registration_is_not_listed_for_supervisor(self):
        resp = self.client.get(reverse('supervisor:users'))

        self.assertEqual(resp.status_code, 200)
        self.assertNotIn(self.pending, list(resp.context['pending']))

    def test_verified_registration_is_listed_for_supervisor(self):
        self.pending.email_verified = True
        self.pending.save(update_fields=['email_verified'])

        resp = self.client.get(reverse('supervisor:users'))

        self.assertEqual(resp.status_code, 200)
        self.assertIn(self.pending, list(resp.context['pending']))

    def test_can_approve_verified_email_registration(self):
        self.pending.email_verified = True
        self.pending.save(update_fields=['email_verified'])

        resp = self.client.post(reverse('supervisor:approve_registration', args=[self.pending.id]))

        self.assertRedirects(resp, reverse('supervisor:users'), fetch_redirect_response=False)
        self.pending.refresh_from_db()
        self.assertTrue(self.pending.is_active)
        self.assertFalse(self.pending.is_pending_approval)
